import pytest
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from ai_reviewer.article import ArticleText, Chapter
from ai_reviewer.splitter import NO_HEADING, ArticleParser, ArticleReadError

parse_article = ArticleParser().parse
article_text = ArticleText().write


def save(doc, tmp_path, name="a.docx"):
    path = tmp_path / name
    doc.save(path)
    return path


def chapters_of(path):
    return [(ch.name, ch.paragraphs) for ch in parse_article(path)]


def test_styles_beat_heuristics(tmp_path):
    doc = Document()
    doc.add_heading("Введение", 1)
    doc.add_paragraph().add_run("Жирная строка").bold = True
    doc.add_paragraph("Текст.")

    assert chapters_of(save(doc, tmp_path)) == [("Введение", ["Жирная строка", "Текст."])]


def test_no_headings(tmp_path):
    doc = Document()
    for text in ["Первый абзац.", "", "Второй абзац.", "   ", "Третий."]:
        doc.add_paragraph(text)

    assert chapters_of(save(doc, tmp_path)) == [(NO_HEADING, ["Первый абзац.", "Второй абзац.", "Третий."])]


def test_chapter_ids(tmp_path):
    doc = Document()
    for title in ['A/B:C*?"<>|D', "CON", "д" * 150, "Итоги", "итоги", "Итоги"]:
        doc.add_heading(title, 1)

    names = [name for name, _ in chapters_of(save(doc, tmp_path))]

    assert names == ['A/B:C*?"<>|D', "CON", "д" * 150, "Итоги", "итоги", "Итоги (2)"]


def test_unreadable_article(tmp_path):
    (tmp_path / "a.docx").write_bytes(b"not a zip")

    for path in (tmp_path / "missing.docx", tmp_path / "a.docx"):
        with pytest.raises(ArticleReadError):
            parse_article(path)


def test_table_and_list(tmp_path):
    doc = Document()
    doc.add_paragraph("До.")
    table = doc.add_table(rows=2, cols=2)
    for cell, text in zip(table._cells, "abcd"):
        cell.text = text
    doc.add_paragraph("После.")
    doc.add_paragraph("один;", style="List Number")
    doc.add_paragraph("два.", style="List Bullet")

    assert chapters_of(save(doc, tmp_path)) == [
        (NO_HEADING, ["До.", "a\tb\nc\td", "После.", "один;", "два."])
    ]


def test_exclusions_and_tracked_changes(tmp_path):
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "HDR"
    doc.sections[0].footer.paragraphs[0].text = "FTR"
    doc.add_paragraph("Формула ")._p.append(parse_xml(
        '<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
        "<m:r><m:t>EQ</m:t></m:r></m:oMath>"
    ))
    edited = doc.add_paragraph("Правка ")._p
    edited.append(parse_xml(f'<w:ins {nsdecls("w")} w:id="1" w:author="a"><w:r><w:t>INS</w:t></w:r></w:ins>'))
    edited.append(parse_xml(
        f'<w:del {nsdecls("w")} w:id="2" w:author="a"><w:r><w:delText>DEL</w:delText></w:r></w:del>'
    ))

    assert chapters_of(save(doc, tmp_path)) == [(NO_HEADING, ["Формула", "Правка INS"])]


def test_article_text_format():
    chapters = [Chapter("A", ["p1", "p2"]), Chapter("B", ["q1"])]
    text = article_text(chapters)

    assert text == (
        "!! chapter: A\n!! paragraph: 1\np1\n\n!! paragraph: 2\np2\n\n!! chapter: B\n!! paragraph: 1\nq1\n"
    )
    assert ArticleText().read(text).paragraph_text("A", 2) == "p2"


def test_article_text_name_verbatim_and_empty_chapter():
    chapters = [
        Chapter(NO_HEADING, ["Г.\xa0А.\xa0Шебанин"]),
        Chapter("A_B", ["a\nb\tc"]),
        Chapter("Пусто"),
        Chapter("C", ["d"]),
    ]

    assert article_text(chapters) == (
        f"!! chapter: {NO_HEADING}\n!! paragraph: 1\nГ.\xa0А.\xa0Шебанин\n\n"
        "!! chapter: A_B\n!! paragraph: 1\na\nb\tc\n\n"
        "!! chapter: Пусто\n\n"
        "!! chapter: C\n!! paragraph: 1\nd\n"
    )


def test_article_text_edges():
    assert article_text([]) == ""
    assert article_text([Chapter("A", ["p"]), Chapter("B")]).endswith("p\n\n!! chapter: B\n")
